import re
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

DRIVE_DIR = Path(
    r"G:\Mi unidad\Borrar - Dualismo filosófico - Eneko\Espiritualismo\Estudios\Euskal Herriko Unibertsitatea\Compartido Eneko\4. Kurtsoa\4. Kurtsoa 1. Lauhilekoa\RKA\rka_6\2. Proiektua"
)
OUTPUT_DOCX = DRIVE_DIR / "Praktika_Txostena_RKA.docx"
LOCAL_DOCX = Path(__file__).resolve().parent / "Praktika_Txostena_RKA.docx"

BASE_DIR = Path(__file__).resolve().parent


def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>'))


def clean_latex(text: str) -> str:
    t = text
    # Replace common formatting
    t = re.sub(r'\\textbf\{([^}]+)\}', r'\1', t)
    t = re.sub(r'\\texttt\{([^}]+)\}', r'\1', t)
    t = re.sub(r'\\emph\{([^}]+)\}', r'\1', t)
    t = re.sub(r'\\egiaztatu\{([^}]+)\}', r'\1', t)
    t = re.sub(r'\\bete\{([^}]+)\}', r'\1', t)
    t = re.sub(r'\\ref\{[^}]+\}', r'[Erreferentzia]', t)
    t = re.sub(r'\\cite\{[^}]+\}', r'[Erref]', t)
    t = re.sub(r'\\label\{[^}]+\}', '', t)
    t = t.replace('~', ' ')
    t = t.replace(r'\,', ' ')
    t = t.replace(r'\_', '_')
    t = t.replace(r'\%', '%')
    t = t.replace(r'\dots', '...')
    t = t.replace('{,}', ',')
    # Math symbols
    t = t.replace(r'\pm', '±')
    t = t.replace(r'\le', '≤')
    t = t.replace(r'\ge', '≥')
    t = t.replace(r'\approx', '≈')
    t = t.replace(r'\pi', 'π')
    t = t.replace(r'\theta', 'θ')
    t = t.replace(r'\kappa', 'κ')
    t = t.replace(r'\omega', 'ω')
    t = t.replace(r'^\circ', '°')
    t = re.sub(r'\$([^$]+)\$', r'\1', t)
    return t.strip()


def add_code_block(doc, code_text):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, "F4F5F7")
    cell.width = Inches(6.5)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(code_text.strip())
    run.font.name = "Consolas"
    run.font.size = Pt(9.5)
    run.font.color.rgb = RGBColor(30, 41, 59)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def build_document():
    doc = Document()
    
    # Configure margins
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(12)
    p_title.paragraph_format.space_after = Pt(4)
    r_title = p_title.add_run("Oztopoak ekiditen eta errepideari jarraitzen")
    r_title.bold = True
    r_title.font.name = "Calibri"
    r_title.font.size = Pt(22)
    r_title.font.color.rgb = RGBColor(15, 23, 42)

    # Subtitle
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(14)
    r_sub = p_sub.add_run("Sentsoreen funtzionamendua lantzen: ROSMASTER X3 robotarekin egindako praktika")
    r_sub.italic = True
    r_sub.font.size = Pt(12)
    r_sub.font.color.rgb = RGBColor(71, 85, 105)

    # Authors & Info box
    p_meta = doc.add_paragraph()
    p_meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_meta.paragraph_format.space_after = Pt(24)
    r_meta = p_meta.add_run(
        "Egileak: Eneko Ruiz, Eneko Pardavila, Omar Ouinas\n"
        "Informatika Ingeniaritzako Gradua · Robotika eta Kontrol Adimendunak (Irakaslea: Elena Lazkano)\n"
        "Euskal Herriko Unibertsitatea (UPV/EHU)"
    )
    r_meta.font.size = Pt(10)
    r_meta.font.color.rgb = RGBColor(100, 116, 139)

    doc.add_heading("Laburpena", level=1)
    laburpena_path = BASE_DIR / "kapituluak" / "laburpena.tex"
    if laburpena_path.exists():
        text = laburpena_path.read_text(encoding="utf-8")
        paras = [p.strip() for p in text.split("\n\n") if p.strip() and not p.strip().startswith("%")]
        for p in paras:
            doc.add_paragraph(clean_latex(p))

    chapters = [
        ("1. Sarrera", "01_sarrera.tex"),
        ("2. Oztopoak ekiditen (LiDAR)", "02_oztopoak.tex"),
        ("3. Errepideari jarraitzen (Kamera)", "03_errepidea.tex"),
        ("4. Ondorioak", "04_ondorioak.tex"),
    ]

    for ch_title, filename in chapters:
        f_path = BASE_DIR / "kapituluak" / filename
        if not f_path.exists():
            continue
        
        doc.add_heading(ch_title, level=1)
        content = f_path.read_text(encoding="utf-8")
        
        # Parse blocks
        parse_chapter_content(doc, content)

    # Save locally and to Drive
    if LOCAL_DOCX.exists():
        try:
            LOCAL_DOCX.unlink()
        except Exception:
            pass
    doc.save(str(LOCAL_DOCX))
    print(f"[OK] Guardado localmente: {LOCAL_DOCX}")
    
    if DRIVE_DIR.exists():
        if OUTPUT_DOCX.exists():
            try:
                OUTPUT_DOCX.unlink()
                print(f"[OK] Eliminada la version anterior en Drive: {OUTPUT_DOCX.name}")
            except Exception as e:
                print(f"[AVISO] No se pudo borrar el archivo anterior (posiblemente abierto), se sobreescribira: {e}")
        doc.save(str(OUTPUT_DOCX))
        print(f"[OK] Guardada la nueva version limpia en Google Drive: {OUTPUT_DOCX}")
    else:
        print(f"[ALERTA] No se encontro la carpeta de Drive: {DRIVE_DIR}")


def parse_chapter_content(doc, content: str):
    # Regex patterns for elements
    lines = content.splitlines()
    in_code = False
    code_lines = []
    in_table = False
    table_lines = []
    
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        
        # Comments
        if line.startswith("%"):
            i += 1
            continue

        # Code block start
        if "\\begin{lstlisting}" in line:
            in_code = True
            code_lines = []
            i += 1
            continue
        elif "\\end{lstlisting}" in line:
            in_code = False
            add_code_block(doc, "\n".join(code_lines))
            i += 1
            continue
        elif in_code:
            code_lines.append(lines[i])
            i += 1
            continue

        # Table block start
        if "\\begin{table}" in line or "\\begin{tabular" in line:
            in_table = True
            table_lines = [line]
            i += 1
            while i < len(lines) and "\\end{table}" not in lines[i] and "\\end{tabular" not in lines[i]:
                table_lines.append(lines[i])
                i += 1
            if i < len(lines):
                table_lines.append(lines[i])
                i += 1
            parse_latex_table(doc, "\n".join(table_lines))
            in_table = False
            continue

        # Headings
        m_sec = re.match(r'\\section\{([^}]+)\}', line)
        if m_sec:
            doc.add_heading(clean_latex(m_sec.group(1)), level=2)
            i += 1
            continue

        m_ssec = re.match(r'\\subsection\{([^}]+)\}', line)
        if m_ssec:
            doc.add_heading(clean_latex(m_ssec.group(1)), level=3)
            i += 1
            continue

        m_para = re.match(r'\\paragraph\{([^}]+)\}', line)
        if m_para:
            p = doc.add_paragraph()
            r = p.add_run(clean_latex(m_para.group(1)) + " ")
            r.bold = True
            # Check if text follows on same line or next
            rest = line[m_para.end():].strip()
            if rest:
                p.add_run(clean_latex(rest))
            i += 1
            continue

        # Itemize / items
        if line.startswith("\\item"):
            item_text = line[len("\\item"):].strip()
            m_bold = re.match(r'\[([^\]]+)\]', item_text)
            if m_bold:
                prefix = m_bold.group(1)
                body = item_text[m_bold.end():].strip()
                p = doc.add_paragraph(style='List Bullet')
                r = p.add_run(clean_latex(prefix) + " ")
                r.bold = True
                p.add_run(clean_latex(body))
            else:
                p = doc.add_paragraph(clean_latex(item_text), style='List Bullet')
            i += 1
            continue

        # Equations
        if line.startswith("\\begin{equation}"):
            eq_lines = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("\\end{equation}"):
                eq_lines.append(lines[i].strip())
                i += 1
            i += 1
            eq_text = clean_latex(" ".join(eq_lines))
            p_eq = doc.add_paragraph()
            p_eq.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r_eq = p_eq.add_run(eq_text)
            r_eq.italic = True
            r_eq.font.size = Pt(11)
            continue

        # Ignore TikZ, figures placeholders, labels
        if "\\begin{tikzpicture}" in line:
            while i < len(lines) and "\\end{tikzpicture}" not in lines[i]:
                i += 1
            i += 1
            continue

        if "\\irudihutsa" in line:
            m_fig = re.search(r'\\irudihutsa\{([^}]+)\}', line)
            if m_fig:
                p = doc.add_paragraph()
                r = p.add_run(f"[IRUDIA: {clean_latex(m_fig.group(1))}]")
                r.italic = True
                r.font.color.rgb = RGBColor(100, 116, 139)
            i += 1
            continue

        if "\\caption{" in line:
            m_cap = re.search(r'\\caption\{([^}]+)\}', line)
            if m_cap:
                p = doc.add_paragraph()
                r = p.add_run(f"Irudia: {clean_latex(m_cap.group(1))}")
                r.font.size = Pt(9.5)
                r.italic = True
                r.font.color.rgb = RGBColor(100, 116, 139)
            i += 1
            continue

        if line in ("\\begin{itemize}", "\\end{itemize}", "\\begin{description}", "\\end{description}", "\\begin{figure}", "\\end{figure}", "\\centering", "\\clearpage"):
            i += 1
            continue

        # Regular paragraph
        if line:
            # Collect full paragraph till empty line or next command
            p_lines = [line]
            i += 1
            while i < len(lines):
                next_l = lines[i].strip()
                if not next_l or next_l.startswith("\\"):
                    break
                p_lines.append(next_l)
                i += 1
            doc.add_paragraph(clean_latex(" ".join(p_lines)))
        else:
            i += 1


def parse_latex_table(doc, table_text: str):
    caption = ""
    m_cap = re.search(r'\\caption\{([^}]+)\}', table_text)
    if m_cap:
        caption = clean_latex(m_cap.group(1))

    # Extract rows between \midrule (or \toprule) and \bottomrule
    rows = []
    for line in table_text.splitlines():
        line = line.strip()
        if not line or any(cmd in line for cmd in ("\\begin", "\\end", "\\toprule", "\\midrule", "\\bottomrule", "\\caption", "\\centering", "\\label")):
            continue
        if "\\multicolumn" in line:
            # footer row
            m_txt = re.search(r'\{@\{\}l\}\{([^}]+)\}', line)
            if m_txt:
                rows.append([clean_latex(m_txt.group(1))])
            continue
        if "\\\\" in line:
            line = line.replace("\\\\", "")
        parts = [clean_latex(c) for c in line.split("&")]
        if any(p.strip() for p in parts):
            rows.append(parts)

    if not rows:
        return

    num_cols = max(len(r) for r in rows)
    tbl = doc.add_table(rows=len(rows), cols=num_cols)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    for r_idx, r_data in enumerate(rows):
        is_header = (r_idx == 0)
        is_colspan = (len(r_data) == 1 and num_cols > 1)
        
        if is_colspan:
            cell = tbl.cell(r_idx, 0)
            cell.text = r_data[0]
            # Merge across row
            for c_idx in range(1, num_cols):
                cell.merge(tbl.cell(r_idx, c_idx))
            cell.paragraphs[0].runs[0].bold = True
            set_cell_background(cell, "F8FAFC")
            continue

        for c_idx, val in enumerate(r_data):
            cell = tbl.cell(r_idx, c_idx)
            cell.text = val
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(3)
            if is_header:
                set_cell_background(cell, "E2E8F0")
                if p.runs:
                    p.runs[0].bold = True
            else:
                set_cell_background(cell, "FFFFFF" if r_idx % 2 == 1 else "F8FAFC")

    if caption:
        p_cap = doc.add_paragraph()
        r = p_cap.add_run(f"Taula: {caption}")
        r.font.size = Pt(9.5)
        r.italic = True
        r.font.color.rgb = RGBColor(100, 116, 139)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)


if __name__ == "__main__":
    build_document()
